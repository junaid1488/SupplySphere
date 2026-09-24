from __future__ import annotations
from pathlib import Path
from typing import Any
import io

import pandas as pd

from dataset_analyzer.models import DatasetSession, FileFormat
from dataset_analyzer.storage import PathSafetyError


class ReaderError(Exception):
    pass


class DatasetReader:
    MAX_PREVIEW_ROWS = 100
    MAX_SAMPLE_ROWS = 1000
    MAX_COLUMNS = 500
    CSV_CHUNK_SIZE = 10_000

    def __init__(self, session: DatasetSession):
        self._session = session
        self._storage_path = session.storage_path
        self._format = session.metadata.file_format
        self._cached_sheets: list[str] | None = None

    @property
    def session(self) -> DatasetSession:
        return self._session

    @property
    def file_format(self) -> FileFormat:
        return self._format

    @property
    def storage_path(self) -> Path:
        return self._storage_path

    def _validate_column_access(self, columns: list[str] | None, available_columns: list[str]) -> list[str] | None:
        if columns is None:
            return None
        if len(columns) > self.MAX_COLUMNS:
            raise ReaderError(f"Too many columns requested: {len(columns)} > {self.MAX_COLUMNS}")
        invalid = [c for c in columns if c not in available_columns]
        if invalid:
            raise ReaderError(f"Invalid column names: {invalid}")
        return columns

    def _validate_row_bounds(self, n_rows: int, max_allowed: int) -> int:
        if n_rows <= 0:
            raise ReaderError("Row count must be positive")
        if n_rows > max_allowed:
            raise ReaderError(f"Row count {n_rows} exceeds maximum allowed {max_allowed}")
        return n_rows

    def get_preview(self, n_rows: int = 50, columns: list[str] | None = None) -> dict[str, Any]:
        n_rows = self._validate_row_bounds(n_rows, self.MAX_PREVIEW_ROWS)

        if self._format == FileFormat.CSV:
            return self._preview_csv(n_rows, columns)
        elif self._format in (FileFormat.XLS, FileFormat.XLSX):
            return self._preview_excel(n_rows, columns)
        else:
            raise ReaderError(f"Unsupported format for preview: {self._format}")

    def _preview_csv(self, n_rows: int, columns: list[str] | None) -> dict[str, Any]:
        try:
            if columns is not None:
                df = pd.read_csv(self._storage_path, nrows=n_rows, usecols=columns)
            else:
                df = pd.read_csv(self._storage_path, nrows=n_rows)
        except pd.errors.EmptyDataError:
            return {"rows": [], "columns": [], "dtypes": {}, "row_count": 0, "truncated": False}
        except ValueError as e:
            if "usecols" in str(e).lower():
                available = pd.read_csv(self._storage_path, nrows=0).columns.tolist()
                raise ReaderError(f"Invalid column names: {columns}") from e
            raise ReaderError(f"Failed to read CSV preview: {e}") from e
        except Exception as e:
            raise ReaderError(f"Failed to read CSV preview: {e}") from e

        return {
            "rows": df.to_dict(orient="records"),
            "columns": list(df.columns),
            "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
            "row_count": len(df),
            "truncated": len(df) >= n_rows,
        }

    def _preview_excel(self, n_rows: int, columns: list[str] | None) -> dict[str, Any]:
        sheet_name = 0
        try:
            if columns is not None:
                df = pd.read_excel(self._storage_path, sheet_name=sheet_name, nrows=n_rows, usecols=columns)
            else:
                df = pd.read_excel(self._storage_path, sheet_name=sheet_name, nrows=n_rows)
        except ValueError as e:
            if "sheet" in str(e).lower():
                raise ReaderError(f"Sheet not found: {sheet_name}") from e
            if "usecols" in str(e).lower():
                available = pd.read_excel(self._storage_path, sheet_name=sheet_name, nrows=0).columns.tolist()
                raise ReaderError(f"Invalid column names: {columns}") from e
            raise ReaderError(f"Failed to read Excel preview: {e}") from e
        except Exception as e:
            raise ReaderError(f"Failed to read Excel preview: {e}") from e

        return {
            "rows": df.to_dict(orient="records"),
            "columns": list(df.columns),
            "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
            "row_count": len(df),
            "truncated": len(df) >= n_rows,
            "sheet_name": self.get_sheet_names()[0] if self.get_sheet_names() else None,
        }

    def get_sample(self, n_rows: int = 100, columns: list[str] | None = None, random_state: int | None = 42) -> dict[str, Any]:
        n_rows = self._validate_row_bounds(n_rows, self.MAX_SAMPLE_ROWS)

        if self._format == FileFormat.CSV:
            return self._sample_csv(n_rows, columns, random_state)
        elif self._format in (FileFormat.XLS, FileFormat.XLSX):
            return self._sample_excel(n_rows, columns, random_state)
        else:
            raise ReaderError(f"Unsupported format for sampling: {self._format}")

    def _sample_csv(self, n_rows: int, columns: list[str] | None, random_state: int | None) -> dict[str, Any]:
        try:
            total_rows = self._count_csv_rows()
            if total_rows == 0:
                return {"rows": [], "columns": [], "dtypes": {}, "row_count": 0, "total_rows": 0}

            if total_rows <= n_rows:
                df = pd.read_csv(self._storage_path, usecols=columns) if columns else pd.read_csv(self._storage_path)
            else:
                df = pd.read_csv(self._storage_path, usecols=columns) if columns else pd.read_csv(self._storage_path)
                df = df.sample(n=n_rows, random_state=random_state)
        except pd.errors.EmptyDataError:
            return {"rows": [], "columns": [], "dtypes": {}, "row_count": 0, "total_rows": 0}
        except ValueError as e:
            if "usecols" in str(e).lower():
                available = pd.read_csv(self._storage_path, nrows=0).columns.tolist()
                raise ReaderError(f"Invalid column names: {columns}") from e
            raise ReaderError(f"Failed to read CSV sample: {e}") from e
        except Exception as e:
            raise ReaderError(f"Failed to read CSV sample: {e}") from e

        return {
            "rows": df.to_dict(orient="records"),
            "columns": list(df.columns),
            "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
            "row_count": len(df),
            "total_rows": total_rows,
        }

    def _count_csv_rows(self) -> int:
        try:
            with open(self._storage_path, "r", encoding="utf-8", errors="ignore") as f:
                return sum(1 for _ in f) - 1
        except Exception:
            return 0

    def _sample_excel(self, n_rows: int, columns: list[str] | None, random_state: int | None) -> dict[str, Any]:
        sheet_name = 0
        try:
            df = pd.read_excel(self._storage_path, sheet_name=sheet_name, usecols=columns) if columns else pd.read_excel(self._storage_path, sheet_name=sheet_name)
        except ValueError as e:
            if "sheet" in str(e).lower():
                raise ReaderError(f"Sheet not found: {sheet_name}") from e
            if "usecols" in str(e).lower():
                available = pd.read_excel(self._storage_path, sheet_name=sheet_name, nrows=0).columns.tolist()
                raise ReaderError(f"Invalid column names: {columns}") from e
            raise ReaderError(f"Failed to read Excel for sampling: {e}") from e
        except Exception as e:
            raise ReaderError(f"Failed to read Excel for sampling: {e}") from e

        total_rows = len(df)
        if total_rows == 0:
            return {"rows": [], "columns": [], "dtypes": {}, "row_count": 0, "total_rows": 0}

        if total_rows <= n_rows:
            sampled = df
        else:
            sampled = df.sample(n=n_rows, random_state=random_state)

        return {
            "rows": sampled.to_dict(orient="records"),
            "columns": list(sampled.columns),
            "dtypes": {col: str(dtype) for col, dtype in sampled.dtypes.items()},
            "row_count": len(sampled),
            "total_rows": total_rows,
            "sheet_name": self.get_sheet_names()[0] if self.get_sheet_names() else None,
        }

    def get_sheet_names(self) -> list[str]:
        if self._cached_sheets is not None:
            return self._cached_sheets

        if self._format not in (FileFormat.XLS, FileFormat.XLSX):
            return []

        try:
            with pd.ExcelFile(self._storage_path) as xls:
                self._cached_sheets = xls.sheet_names
            return self._cached_sheets
        except Exception as e:
            raise ReaderError(f"Failed to read Excel sheet names: {e}") from e

    def read_sheet(self, sheet_name: str | int = 0, n_rows: int | None = None, columns: list[str] | None = None) -> dict[str, Any]:
        if self._format not in (FileFormat.XLS, FileFormat.XLSX):
            raise ReaderError("Sheet reading only supported for Excel formats")

        sheet_names = self.get_sheet_names()
        if isinstance(sheet_name, str):
            if sheet_name not in sheet_names:
                raise ReaderError(f"Sheet not found: {sheet_name}")
        elif isinstance(sheet_name, int):
            if sheet_name < 0 or sheet_name >= len(sheet_names):
                raise ReaderError(f"Sheet index out of range: {sheet_name}")
            sheet_name = sheet_names[sheet_name]

        if n_rows is not None:
            n_rows = self._validate_row_bounds(n_rows, self.MAX_PREVIEW_ROWS * 10)

        try:
            if n_rows is not None:
                df = pd.read_excel(self._storage_path, sheet_name=sheet_name, nrows=n_rows, usecols=columns) if columns else pd.read_excel(self._storage_path, sheet_name=sheet_name, nrows=n_rows)
            else:
                df = pd.read_excel(self._storage_path, sheet_name=sheet_name, usecols=columns) if columns else pd.read_excel(self._storage_path, sheet_name=sheet_name)
        except ValueError as e:
            if "usecols" in str(e).lower():
                available = pd.read_excel(self._storage_path, sheet_name=sheet_name, nrows=0).columns.tolist()
                raise ReaderError(f"Invalid column names: {columns}") from e
            raise ReaderError(f"Failed to read Excel sheet: {e}") from e
        except Exception as e:
            raise ReaderError(f"Failed to read Excel sheet: {e}") from e

        return {
            "rows": df.to_dict(orient="records"),
            "columns": list(df.columns),
            "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
            "row_count": len(df),
            "sheet_name": sheet_name,
            "truncated": n_rows is not None and len(df) >= n_rows,
        }

    def read_csv_chunked(
        self,
        chunk_size: int = 10_000,
        columns: list[str] | None = None,
        max_chunks: int | None = None,
    ) -> list[dict[str, Any]]:
        if self._format != FileFormat.CSV:
            raise ReaderError("Chunked reading only supported for CSV format")

        chunk_size = min(chunk_size, self.CSV_CHUNK_SIZE)
        chunks = []

        try:
            reader = pd.read_csv(self._storage_path, chunksize=chunk_size, usecols=columns) if columns else pd.read_csv(self._storage_path, chunksize=chunk_size)
        except ValueError as e:
            if "usecols" in str(e).lower():
                available = pd.read_csv(self._storage_path, nrows=0).columns.tolist()
                raise ReaderError(f"Invalid column names: {columns}") from e
            raise ReaderError(f"Failed to create CSV chunked reader: {e}") from e
        except Exception as e:
            raise ReaderError(f"Failed to create CSV chunked reader: {e}") from e

        for i, chunk in enumerate(reader):
            if max_chunks is not None and i >= max_chunks:
                break

            chunks.append({
                "rows": chunk.to_dict(orient="records"),
                "columns": list(chunk.columns),
                "dtypes": {col: str(dtype) for col, dtype in chunk.dtypes.items()},
                "row_count": len(chunk),
                "chunk_index": i,
            })

        return chunks

    def get_column_info(self) -> dict[str, Any]:
        if self._format == FileFormat.CSV:
            try:
                df = pd.read_csv(self._storage_path, nrows=0)
            except pd.errors.EmptyDataError:
                return {"columns": [], "dtypes": {}, "column_count": 0}
            except Exception as e:
                raise ReaderError(f"Failed to read CSV columns: {e}") from e
        elif self._format in (FileFormat.XLS, FileFormat.XLSX):
            try:
                df = pd.read_excel(self._storage_path, sheet_name=0, nrows=0)
            except Exception as e:
                raise ReaderError(f"Failed to read Excel columns: {e}") from e
        else:
            raise ReaderError(f"Unsupported format: {self._format}")

        return {
            "columns": list(df.columns),
            "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
            "column_count": len(df.columns),
        }

    def get_row_count(self) -> int:
        if self._format == FileFormat.CSV:
            return self._count_csv_rows()
        elif self._format in (FileFormat.XLS, FileFormat.XLSX):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(self._storage_path, read_only=True, data_only=True)
                try:
                    ws = wb.active
                    if ws is None:
                        return 0
                    max_row = ws.max_row
                    return max(int(max_row) - 1, 0) if max_row else 0
                finally:
                    wb.close()
            except Exception:
                try:
                    df = pd.read_excel(self._storage_path, sheet_name=0, nrows=0)
                    return 0
                except Exception:
                    return 0
        return 0


def create_reader(session: DatasetSession) -> DatasetReader:
    return DatasetReader(session)