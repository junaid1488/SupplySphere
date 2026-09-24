from __future__ import annotations
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
import math
import re

import numpy as np
import pandas as pd

from dataset_analyzer.models import (
    ColumnProfile,
    DatasetProfile,
    FileFormat,
    InferredType,
)


@dataclass
class ProfilingConfig:
    max_rows_for_full_profile: int = 100_000
    max_unique_for_categorical_summary: int = 50
    sample_size_for_stats: int = 10_000
    chunk_size: int = 10_000
    max_sample_values: int = 10


class DatasetProfiler:
    def __init__(self, session_path: Path, file_format: FileFormat):
        self._session_path = session_path
        self._file_format = file_format
        self._config = ProfilingConfig()

    def profile(self) -> DatasetProfile:
        if self._file_format == FileFormat.CSV:
            return self._profile_csv()
        elif self._file_format in (FileFormat.XLS, FileFormat.XLSX):
            return self._profile_excel()
        else:
            raise ValueError(f"Unsupported format: {self._file_format}")

    def _profile_csv(self) -> DatasetProfile:
        total_rows = self._count_csv_rows()
        if total_rows == 0:
            return self._empty_profile()

        column_info = self._get_column_info()
        columns = column_info["columns"]

        if total_rows <= self._config.max_rows_for_full_profile:
            df = pd.read_csv(self._session_path)
            return self._compute_full_profile(df, total_rows, columns)
        else:
            return self._compute_sampled_profile(total_rows, columns)

    def _profile_excel(self) -> DatasetProfile:
        max_rows = self._config.max_rows_for_full_profile
        df = pd.read_excel(self._session_path, sheet_name=0, nrows=max_rows + 1)
        read_rows = len(df)
        columns = list(df.columns)

        if read_rows == 0:
            return self._empty_profile()

        if read_rows <= max_rows:
            return self._compute_full_profile(df, read_rows, columns)

        total_rows = self._count_excel_rows()
        return self._compute_sampled_profile_from_df(df, total_rows, columns)

    def _count_excel_rows(self) -> int:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(self._session_path, read_only=True, data_only=True)
            try:
                ws = wb.active
                if ws is None:
                    return self._config.max_rows_for_full_profile + 1
                max_row = ws.max_row
                if not max_row:
                    return self._config.max_rows_for_full_profile + 1
                return max(int(max_row) - 1, 0)
            finally:
                wb.close()
        except Exception:
            return self._config.max_rows_for_full_profile + 1

    def _count_csv_rows(self) -> int:
        try:
            with open(self._session_path, "r", encoding="utf-8", errors="ignore") as f:
                return sum(1 for _ in f) - 1
        except Exception:
            return 0

    def _get_column_info(self) -> dict[str, Any]:
        try:
            df = pd.read_csv(self._session_path, nrows=0)
        except pd.errors.EmptyDataError:
            return {"columns": [], "dtypes": {}, "column_count": 0}
        except Exception:
            return {"columns": [], "dtypes": {}, "column_count": 0}

        return {
            "columns": list(df.columns),
            "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
            "column_count": len(df.columns),
        }

    def _empty_profile(self) -> DatasetProfile:
        return DatasetProfile(
            dataset_id=self._session_path.parent.name,
            row_count=0,
            column_count=0,
            column_names=[],
            column_profiles={},
            duplicate_row_count=0,
            numeric_columns=[],
            categorical_columns=[],
            datetime_columns=[],
            boolean_columns=[],
            id_key_candidates=[],
            measure_candidates=[],
            geographic_candidates=[],
            profiled_at=datetime.utcnow(),
            limitations=["Dataset is empty"],
        )

    def _compute_full_profile(
        self, df: pd.DataFrame, total_rows: int, columns: list[str]
    ) -> DatasetProfile:
        column_profiles = {}
        numeric_columns = []
        categorical_columns = []
        datetime_columns = []
        boolean_columns = []
        id_key_candidates = []
        measure_candidates = []
        geographic_candidates = []
        limitations = []

        for col in columns:
            profile = self._profile_column(df[col], col, total_rows)
            column_profiles[col] = profile

            if profile.inferred_type == InferredType.NUMERIC or profile.inferred_type in (
                InferredType.INTEGER,
                InferredType.FLOAT,
            ):
                numeric_columns.append(col)
                if profile.is_measure_candidate:
                    measure_candidates.append(col)
            elif profile.inferred_type in (InferredType.DATETIME, InferredType.DATE, InferredType.TIME):
                datetime_columns.append(col)
            elif profile.inferred_type == InferredType.BOOLEAN:
                boolean_columns.append(col)
            else:
                categorical_columns.append(col)

            if profile.is_id_candidate:
                id_key_candidates.append(col)
            if profile.is_geographic_candidate:
                geographic_candidates.append(col)

        duplicate_row_count = int(df.duplicated().sum())

        return DatasetProfile(
            dataset_id=self._session_path.parent.name,
            row_count=total_rows,
            column_count=len(columns),
            column_names=columns,
            column_profiles=column_profiles,
            duplicate_row_count=duplicate_row_count,
            numeric_columns=numeric_columns,
            categorical_columns=categorical_columns,
            datetime_columns=datetime_columns,
            boolean_columns=boolean_columns,
            id_key_candidates=id_key_candidates,
            measure_candidates=measure_candidates,
            geographic_candidates=geographic_candidates,
            profiled_at=datetime.utcnow(),
            limitations=limitations,
        )

    def _compute_sampled_profile(self, total_rows: int, columns: list[str]) -> DatasetProfile:
        sample_size = min(self._config.sample_size_for_stats, total_rows)
        df_sample = pd.read_csv(self._session_path, nrows=sample_size)

        column_profiles = {}
        numeric_columns = []
        categorical_columns = []
        datetime_columns = []
        boolean_columns = []
        id_key_candidates = []
        measure_candidates = []
        geographic_candidates = []
        limitations = [f"Profile computed on sample of {sample_size} rows (total: {total_rows})"]

        for col in columns:
            profile = self._profile_column(df_sample[col], col, total_rows, is_sample=True)
            column_profiles[col] = profile

            if profile.inferred_type == InferredType.NUMERIC or profile.inferred_type in (
                InferredType.INTEGER,
                InferredType.FLOAT,
            ):
                numeric_columns.append(col)
                if profile.is_measure_candidate:
                    measure_candidates.append(col)
            elif profile.inferred_type in (InferredType.DATETIME, InferredType.DATE, InferredType.TIME):
                datetime_columns.append(col)
            elif profile.inferred_type == InferredType.BOOLEAN:
                boolean_columns.append(col)
            else:
                categorical_columns.append(col)

            if profile.is_id_candidate:
                id_key_candidates.append(col)
            if profile.is_geographic_candidate:
                geographic_candidates.append(col)

        duplicate_row_count = -1
        limitations.append("Duplicate row count not computed on sampled data")

        return DatasetProfile(
            dataset_id=self._session_path.parent.name,
            row_count=total_rows,
            column_count=len(columns),
            column_names=columns,
            column_profiles=column_profiles,
            duplicate_row_count=duplicate_row_count,
            numeric_columns=numeric_columns,
            categorical_columns=categorical_columns,
            datetime_columns=datetime_columns,
            boolean_columns=boolean_columns,
            id_key_candidates=id_key_candidates,
            measure_candidates=measure_candidates,
            geographic_candidates=geographic_candidates,
            profiled_at=datetime.utcnow(),
            limitations=limitations,
        )

    def _compute_sampled_profile_from_df(
        self, df: pd.DataFrame, total_rows: int, columns: list[str]
    ) -> DatasetProfile:
        sample_size = min(self._config.sample_size_for_stats, total_rows)
        df_sample = df.sample(n=sample_size, random_state=42) if total_rows > sample_size else df

        column_profiles = {}
        numeric_columns = []
        categorical_columns = []
        datetime_columns = []
        boolean_columns = []
        id_key_candidates = []
        measure_candidates = []
        geographic_candidates = []
        limitations = [f"Profile computed on sample of {sample_size} rows (total: {total_rows})"]

        for col in columns:
            profile = self._profile_column(df_sample[col], col, total_rows, is_sample=True)
            column_profiles[col] = profile

            if profile.inferred_type == InferredType.NUMERIC or profile.inferred_type in (
                InferredType.INTEGER,
                InferredType.FLOAT,
            ):
                numeric_columns.append(col)
                if profile.is_measure_candidate:
                    measure_candidates.append(col)
            elif profile.inferred_type in (InferredType.DATETIME, InferredType.DATE, InferredType.TIME):
                datetime_columns.append(col)
            elif profile.inferred_type == InferredType.BOOLEAN:
                boolean_columns.append(col)
            else:
                categorical_columns.append(col)

            if profile.is_id_candidate:
                id_key_candidates.append(col)
            if profile.is_geographic_candidate:
                geographic_candidates.append(col)

        duplicate_row_count = int(df.duplicated().sum()) if total_rows <= self._config.max_rows_for_full_profile else -1
        if duplicate_row_count == -1:
            limitations.append("Duplicate row count not computed on sampled data")

        return DatasetProfile(
            dataset_id=self._session_path.parent.name,
            row_count=total_rows,
            column_count=len(columns),
            column_names=columns,
            column_profiles=column_profiles,
            duplicate_row_count=duplicate_row_count,
            numeric_columns=numeric_columns,
            categorical_columns=categorical_columns,
            datetime_columns=datetime_columns,
            boolean_columns=boolean_columns,
            id_key_candidates=id_key_candidates,
            measure_candidates=measure_candidates,
            geographic_candidates=geographic_candidates,
            profiled_at=datetime.utcnow(),
            limitations=limitations,
        )

    def _profile_column(
        self, series: pd.Series, col_name: str, total_rows: int, is_sample: bool = False
    ) -> ColumnProfile:
        pandas_dtype = str(series.dtype)
        null_count = int(series.isna().sum())
        null_percentage = (null_count / total_rows * 100) if total_rows > 0 else 0.0
        unique_count = int(series.nunique())
        duplicate_count = int(len(series) - unique_count) if not is_sample else -1

        inferred_type = self._infer_type(series)
        sample_values = self._get_sample_values(series)

        numeric_stats = None
        categorical_summary = None
        datetime_range = None
        boolean_distribution = None

        is_id_candidate = False
        is_geographic_candidate = False
        is_measure_candidate = False

        if inferred_type in (InferredType.INTEGER, InferredType.FLOAT, InferredType.NUMERIC):
            numeric_stats = self._compute_numeric_stats(series)
            is_measure_candidate = self._is_measure_candidate(col_name, series, numeric_stats)
            is_id_candidate = self._is_id_candidate(col_name, series, unique_count, total_rows)
            if is_id_candidate:
                is_measure_candidate = False
        elif inferred_type in (InferredType.DATETIME, InferredType.DATE, InferredType.TIME):
            datetime_range = self._compute_datetime_range(series)
        elif inferred_type == InferredType.BOOLEAN:
            boolean_distribution = self._compute_boolean_distribution(series)
        else:
            categorical_summary = self._compute_categorical_summary(series)
            is_id_candidate = self._is_id_candidate(col_name, series, unique_count, total_rows)

        is_geographic_candidate = self._is_geographic_candidate(col_name, series)

        return ColumnProfile(
            name=col_name,
            inferred_type=inferred_type,
            pandas_dtype=pandas_dtype,
            null_count=null_count,
            null_percentage=round(null_percentage, 2),
            unique_count=unique_count,
            duplicate_count=duplicate_count if not is_sample else -1,
            sample_values=sample_values,
            numeric_stats=numeric_stats,
            categorical_summary=categorical_summary,
            datetime_range=datetime_range,
            boolean_distribution=boolean_distribution,
            is_id_candidate=is_id_candidate,
            is_geographic_candidate=is_geographic_candidate,
            is_measure_candidate=is_measure_candidate,
        )

    def _infer_type(self, series: pd.Series) -> InferredType:
        non_null = series.dropna()
        if len(non_null) == 0:
            return InferredType.UNKNOWN

        dtype = series.dtype

        if pd.api.types.is_bool_dtype(dtype):
            return InferredType.BOOLEAN

        if pd.api.types.is_datetime64_any_dtype(dtype):
            return InferredType.DATETIME

        if pd.api.types.is_integer_dtype(dtype):
            return InferredType.INTEGER

        if pd.api.types.is_float_dtype(dtype):
            return InferredType.FLOAT

        if pd.api.types.is_object_dtype(dtype) or pd.api.types.is_string_dtype(dtype):
            sample = non_null.head(100)
            if self._looks_like_boolean(sample):
                return InferredType.BOOLEAN
            if self._looks_like_datetime(sample):
                return InferredType.DATETIME
            if self._looks_like_numeric(sample):
                return InferredType.NUMERIC
            return InferredType.STRING

        return InferredType.UNKNOWN

    def _looks_like_boolean(self, series: pd.Series) -> bool:
        unique_vals = set(str(v).lower().strip() for v in series.unique()[:20])
        bool_patterns = [
            {"true", "false"},
            {"t", "f"},
            {"yes", "no"},
            {"y", "n"},
            {"1", "0"},
            {"on", "off"},
        ]
        return any(unique_vals.issubset(pattern) for pattern in bool_patterns)

    def _looks_like_datetime(self, series: pd.Series) -> bool:
        sample = series.head(50)
        non_null = sample.dropna()
        if len(non_null) == 0:
            return False

        # Reject columns that look like numeric IDs/codes
        str_vals = non_null.astype(str)
        has_alpha = any(c.isalpha() for v in str_vals for c in str(v))
        pure_numeric_count = 0
        for v in str_vals:
            v_stripped = str(v).strip()
            if v_stripped.lstrip('-').replace('.', '', 1).isdigit():
                pure_numeric_count += 1
        if pure_numeric_count / len(str_vals) > 0.5:
            return False

        # Reject if values contain patterns typical of IDs/codes (letters mixed with digits)
        id_pattern_count = 0
        for v in str_vals:
            v_stripped = str(v).strip()
            if len(v_stripped) > 3 and re.search(r'[A-Za-z]\d|\d[A-Za-z]', v_stripped):
                id_pattern_count += 1
        if id_pattern_count / len(str_vals) > 0.3:
            return False

        try:
            parsed = pd.to_datetime(non_null, errors="coerce")
            parse_rate = parsed.notna().mean()
            if parse_rate <= 0.8:
                return False

            # Additional validation: parsed dates should have reasonable year range
            valid_dates = parsed.dropna()
            if len(valid_dates) == 0:
                return False
            years = valid_dates.dt.year
            if years.min() < 1900 or years.max() > 2100:
                return False

            return True
        except Exception:
            return False

    def _looks_like_numeric(self, series: pd.Series) -> bool:
        sample = series.head(50)
        try:
            parsed = pd.to_numeric(sample, errors="coerce")
            return parsed.notna().mean() > 0.8
        except Exception:
            return False

    def _get_sample_values(self, series: pd.Series) -> list[Any]:
        non_null = series.dropna()
        if len(non_null) == 0:
            return []
        sample_size = min(self._config.max_sample_values, len(non_null))
        if sample_size >= len(non_null):
            return non_null.head(sample_size).tolist()
        return non_null.sample(n=sample_size, random_state=42).tolist()

    def _compute_numeric_stats(self, series: pd.Series) -> dict[str, float]:
        non_null = series.dropna()
        if len(non_null) == 0:
            return {}

        try:
            numeric_series = pd.to_numeric(non_null, errors="coerce").dropna()
            if len(numeric_series) == 0:
                return {}

            return {
                "min": float(numeric_series.min()),
                "max": float(numeric_series.max()),
                "mean": float(numeric_series.mean()),
                "std": float(numeric_series.std()) if len(numeric_series) > 1 else 0.0,
                "median": float(numeric_series.median()),
                "q25": float(numeric_series.quantile(0.25)),
                "q75": float(numeric_series.quantile(0.75)),
                "sum": float(numeric_series.sum()),
                "count": int(len(numeric_series)),
            }
        except Exception:
            return {}

    def _compute_categorical_summary(self, series: pd.Series) -> dict[str, int]:
        non_null = series.dropna()
        if len(non_null) == 0:
            return {}

        value_counts = non_null.value_counts()
        max_unique = self._config.max_unique_for_categorical_summary
        if len(value_counts) > max_unique:
            top_values = value_counts.head(max_unique)
            return {str(k): int(v) for k, v in top_values.items()}
        return {str(k): int(v) for k, v in value_counts.items()}

    def _compute_datetime_range(self, series: pd.Series) -> dict[str, str]:
        non_null = series.dropna()
        if len(non_null) == 0:
            return {}

        try:
            parsed = pd.to_datetime(non_null, errors="coerce").dropna()
            if len(parsed) == 0:
                return {}

            return {
                "min": parsed.min().isoformat(),
                "max": parsed.max().isoformat(),
            }
        except Exception:
            return {}

    def _compute_boolean_distribution(self, series: pd.Series) -> dict[str, int]:
        non_null = series.dropna()
        if len(non_null) == 0:
            return {}

        value_counts = non_null.value_counts()
        return {str(k): int(v) for k, v in value_counts.items()}

    def _is_id_candidate(self, col_name: str, series: pd.Series, unique_count: int, total_rows: int) -> bool:
        if total_rows == 0:
            return False

        name_lower = col_name.lower()
        id_keywords = ["id", "key", "pk", "uuid", "guid", "code", "number", "num"]
        if any(keyword in name_lower for keyword in id_keywords):
            uniqueness_ratio = unique_count / total_rows
            return uniqueness_ratio > 0.95

        uniqueness_ratio = unique_count / total_rows
        return uniqueness_ratio > 0.99

    def _is_measure_candidate(self, col_name: str, series: pd.Series, numeric_stats: dict[str, float] | None) -> bool:
        if not numeric_stats:
            return False

        name_lower = col_name.lower()
        measure_keywords = [
            "price", "cost", "amount", "value", "revenue", "sales", "total",
            "quantity", "qty", "count", "sum", "avg", "average", "mean",
            "weight", "volume", "size", "length", "width", "height",
            "rate", "ratio", "percentage", "percent", "score", "metric",
            "profit", "margin", "income", "expense", "budget", "forecast",
        ]
        id_keywords = ["id", "key", "pk", "uuid", "guid", "code", "number", "num"]

        has_explicit_measure = any(keyword in name_lower for keyword in measure_keywords)
        has_id_pattern = any(keyword in name_lower for keyword in id_keywords)

        if has_explicit_measure:
            return True

        if has_id_pattern:
            return False

        if numeric_stats.get("min", 0) >= 0 and numeric_stats.get("std", 0) > 0:
            return True

        return False

    def _is_geographic_candidate(self, col_name: str, series: pd.Series) -> bool:
        name_lower = col_name.lower()
        # Full-word matches
        full_geo_keywords = [
            "latitude", "longitude", "coordinate", "coordinates",
            "geo", "location",
        ]
        if any(keyword in name_lower for keyword in full_geo_keywords):
            return True

        # Word-boundary matches for short keywords
        if re.search(r'(?<![a-z])lat(?![a-z])', name_lower):
            return True
        if re.search(r'(?<![a-z])lon(?![a-z])', name_lower):
            return True
        if re.search(r'(?<![a-z])lng(?![a-z])', name_lower):
            return True
        if re.search(r'(?<![a-z])loc(?![a-z])', name_lower):
            return True

        non_null = series.dropna()
        if len(non_null) == 0:
            return False

        if pd.api.types.is_numeric_dtype(series):
            sample = non_null.head(100)
            try:
                numeric = pd.to_numeric(sample, errors="coerce").dropna()
                if len(numeric) > 0:
                    if re.search(r'(?<![a-z])lat(?![a-z])', name_lower):
                        return numeric.between(-90, 90).all()
                    if re.search(r'(?<![a-z])lon(?![a-z])', name_lower) or re.search(r'(?<![a-z])lng(?![a-z])', name_lower):
                        return numeric.between(-180, 180).all()
            except Exception:
                pass

        return False


def create_profiler(session_path: Path, file_format: FileFormat) -> DatasetProfiler:
    return DatasetProfiler(session_path, file_format)